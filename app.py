import os

import io

import base64

import sqlite3

import tempfile

import requests

from flask import Flask, request

from openai import OpenAI

from pptx import Presentation

from docx import Document

from reportlab.pdfgen import canvas

app = Flask(__name__)

# =========================

# SETTINGS

# =========================

TELEGRAM_TOKEN = os.environ["TELEGRAM_TOKEN"]

OPENAI_API_KEY = os.environ["OPENAI_API_KEY"]

client = OpenAI(api_key=OPENAI_API_KEY)

TELEGRAM_API = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}"

DB = "student_bot.db"

# =========================

# DATABASE

# =========================

def init_db():

    conn = sqlite3.connect(DB)

    cur = conn.cursor()

    cur.execute("""

        CREATE TABLE IF NOT EXISTS users (

            user_id INTEGER PRIMARY KEY,

            first_name TEXT,

            balance INTEGER DEFAULT 10

        )

    """)

    conn.commit()

    conn.close()

def add_user(user_id, first_name):

    conn = sqlite3.connect(DB)

    cur = conn.cursor()

    cur.execute(

        "INSERT OR IGNORE INTO users (user_id, first_name) VALUES (?, ?)",

        (user_id, first_name)

    )

    conn.commit()

    conn.close()

def get_balance(user_id):

    conn = sqlite3.connect(DB)

    cur = conn.cursor()

    cur.execute(

        "SELECT balance FROM users WHERE user_id = ?",

        (user_id,)

    )

    row = cur.fetchone()

    conn.close()

    return row[0] if row else 0

def use_credit(user_id):

    conn = sqlite3.connect(DB)

    cur = conn.cursor()

    cur.execute(

        "SELECT balance FROM users WHERE user_id = ?",

        (user_id,)

    )

    row = cur.fetchone()

    if not row or row[0] <= 0:

        conn.close()

        return False

    cur.execute(

        "UPDATE users SET balance = balance - 1 WHERE user_id = ?",

        (user_id,)

    )

    conn.commit()

    conn.close()

    return True

# =========================

# TELEGRAM

# =========================

def send_message(chat_id, text):

    requests.post(

        f"{TELEGRAM_API}/sendMessage",

        json={

            "chat_id": chat_id,

            "text": text

        },

        timeout=30

    )

def send_document(chat_id, filepath, caption=""):

    with open(filepath, "rb") as f:

        requests.post(

            f"{TELEGRAM_API}/sendDocument",

            data={

                "chat_id": chat_id,

                "caption": caption

            },

            files={

                "document": f

            },

            timeout=60

        )

def get_file_bytes(file_id):

    result = requests.get(

        f"{TELEGRAM_API}/getFile",

        params={"file_id": file_id},

        timeout=30

    ).json()

    file_path = result["result"]["file_path"]

    data = requests.get(

        f"https://api.telegram.org/file/bot{TELEGRAM_TOKEN}/{file_path}",

        timeout=60

    )

    return data.content

# =========================

# MAIN MENU

# =========================

def main_menu():

    return {

        "inline_keyboard": [

            [

                {"text": "🤖 AI Kömekçi", "callback_data": "ai"},

                {"text": "📷 Suratdan mesele", "callback_data": "image"}

            ],

            [

                {"text": "📊 Prezentasiýa", "callback_data": "ppt"},

                {"text": "📝 Referat / Doklad", "callback_data": "referat"}

            ],

            [

                {"text": "📄 Word", "callback_data": "word"},

                {"text": "📕 PDF", "callback_data": "pdf"}

            ],

            [

                {"text": "🌐 Terjime", "callback_data": "translate"},

                {"text": "💰 Balans", "callback_data": "balance"}

            ],

            [

                {"text": "👤 Profil", "callback_data": "profile"},

                {"text": "🆘 Kömek", "callback_data": "help"}

            ]

        ]

    }

def show_menu(chat_id):

    text = (

        "🎓 STUDENT BOT\n\n"

        "Salam! 👋\n"

        "Men seniň uniwersitet okuw kömekçiň.\n\n"

        "Näme gerek bolsa saýla:"

    )

    requests.post(

        f"{TELEGRAM_API}/sendMessage",

        json={

            "chat_id": chat_id,

            "text": text,

            "reply_markup": main_menu()

        },

        timeout=30

    )

# =========================

# AI

# =========================

def ask_ai(text):

    response = client.responses.create(

        model="gpt-5",

        input=[

            {

                "role": "system",

                "content": (

                    "Sen Student Bot atly professional university AI assistant. "

                    "Türkmençe, rusça we iňlisçe jogap ber. "

                    "Matematika, fizika, himiýa, programmirleme, ykdysadyýet "

                    "we beýleki uniwersitet derslerinde kömek et. "

                    "Meseleleri ädimme-ädim we düşnükli düşündir."

                )

            },

            {

                "role": "user",

                "content": text

            }

        ]

    )

    return response.output_text

# =========================

# IMAGE SOLVING

# =========================

def solve_image(image_bytes):

    encoded = base64.b64encode(image_bytes).decode("utf-8")

    response = client.responses.create(

        model="gpt-5",

        input=[

            {

                "role": "system",

                "content": (

                    "Suratdaky uniwersitet meselesini oka. "

                    "Mesele matematika, fizika, himiýa ýa-da başga ders bolsa "

                    "çözgüdini ädimme-ädim düşündir."

                )

            },

            {

                "role": "user",

                "content": [

                    {

                        "type": "input_text",

                        "text": "Suratdaky meseläni çöz we düşündir."

                    },

                    {

                        "type": "input_image",

                        "image_url": f"data:image/jpeg;base64,{encoded}"

                    }

                ]

            }

        ]

    )

    return response.output_text

# =========================

# PRESENTATION

# =========================

def create_presentation(topic):

    answer = ask_ai(

        f"""

        "{topic}" temasy boýunça uniwersitet derejesinde

        8 slaýdlyk prezentasiýa üçin mazmun taýýarla.

        Her slaýdy:

        SLIDE 1:

        TITLE:

        CONTENT:

        görnüşinde ýaz.

        """

    )

    prs = Presentation()

    slides = answer.split("SLIDE ")

    for block in slides[1:]:

        lines = block.strip().splitlines()

        title = "Student Bot"

        content = ""

        for line in lines:

            if line.startswith("TITLE:"):

                title = line.replace("TITLE:", "").strip()

            elif line.startswith("CONTENT:"):

                content = line.replace("CONTENT:", "").strip()

            elif line.strip():

                content += "\n" + line.strip()

        slide = prs.slides.add_slide(

            prs.slide_layouts[1]

        )

        slide.shapes.title.text = title

        slide.placeholders[1].text = content

    filename = tempfile.mktemp(suffix=".pptx")

    prs.save(filename)

    return filename

# =========================

# WORD

# =========================

def create_word(topic):

    text = ask_ai(

        f"""

        "{topic}" barada uniwersitet üçin gowy gurluşly referat taýýarla.

        Giriş, esasy bölüm, netije we peýdalanylan çeşmeler bölümleri bolsun.

        """

    )

    doc = Document()

    doc.add_heading(topic, 0)

    for paragraph in text.split("\n"):

        if paragraph.strip():

            doc.add_paragraph(paragraph.strip())

    filename = tempfile.mktemp(suffix=".docx")

    doc.save(filename)

    return filename

# =========================

# PDF

# =========================

def create_pdf(topic):

    text = ask_ai(

        f"""

        "{topic}" barada uniwersitet üçin referat taýýarla.

        Giriş, esasy bölüm, netije we çeşmeler bolsun.

        """

    )

    filename = tempfile.mktemp(suffix=".pdf")

    c = canvas.Canvas(filename)

    width, height = 595, 842

    y = height - 50

    c.setFont("Helvetica", 11)

    for paragraph in text.split("\n"):

        if not paragraph.strip():

            y -= 15

            continue

        words = paragraph.split()

        line = ""

        for word in words:

            test = line + " " + word

            if len(test) > 90:

                c.drawString(40, y, line)

                y -= 15

                line = word

                if y < 50:

                    c.showPage()

                    c.setFont("Helvetica", 11)

                    y = height - 50

            else:

                line = test.strip()

        if line:

            c.drawString(40, y, line)

            y -= 15

    c.save()

    return filename

# =========================

# TELEGRAM WEBHOOK

# =========================

@app.route("/", methods=["GET"])

def home():

    return "Student Bot is running! 🎓"

@app.route("/webhook", methods=["POST"])

def webhook():

    data = request.get_json(silent=True) or {}

    # =====================

    # CALLBACK BUTTON

    # =====================

    if "callback_query" in data:

        callback = data["callback_query"]

        chat_id = callback["message"]["chat"]["id"]

        user_id = callback["from"]["id"]

        action = callback["data"]

        if action == "balance":

            balance = get_balance(user_id)

            send_message(

                chat_id,

                f"💰 Balansyň:\n\n⭐ {balance} kredit"

            )

        elif action == "profile":

            balance = get_balance(user_id)

            send_message(

                chat_id,

                f"👤 PROFIL\n\n"

                f"ID: {user_id}\n"

                f"💰 Balans: {balance} kredit"

            )

        elif action == "help":

            send_message(

                chat_id,

                "🆘 KÖMEK\n\n"

                "🤖 AI — soraglaryňy çözýär\n"

                "📷 Surat — suratdaky meseläni çözýär\n"

                "📊 PPT — prezentasiýa taýýarlaýar\n"

                "📝 Referat — referat/doklad ýazýar\n"

                "📄 Word — Word faýly döredýär\n"

                "📕 PDF — PDF döredýär\n"

                "🌐 Terjime — tekst terjime edýär\n\n"

                "Başlamak üçin /start ýaz."

            )

        elif action == "ai":

            send_message(

                chat_id,

                "🤖 AI Kömekçi\n\n"

                "Soragyňy şu görnüşde ýaz:\n\n"

                "Mysal:\n"

                "2x + 5 = 15 meseläni çöz."

            )

        elif action == "image":

            send_message(

                chat_id,

                "📷 Meseläniň suratyny şu ýere iber."

            )

        elif action == "ppt":

            send_message(

                chat_id,

                "📊 Prezentasiýa döretmek üçin:\n\n"

                "/ppt Kompýuter torlarynyň görnüşleri"

            )

        elif action == "referat":

            send_message(

                chat_id,

                "📝 Referat döretmek üçin:\n\n"

                "/referat Emeli intellekt"

            )

        elif action == "word":

            send_message(

                chat_id,

                "📄 Word faýly üçin:\n\n"

                "/word Marketing barada referat"

            )

        elif action == "pdf":

            send_message(

                chat_id,

                "📕 PDF üçin:\n\n"

                "/pdf Emeli intellektiň ösüşi"

            )

        elif action == "translate":

            send_message(

                chat_id,

                "🌐 Terjime etmek üçin:\n\n"

                "/translate Türkmençeden rusça: Salam, nähili?"

            )

        requests.post(

            f"{TELEGRAM_API}/answerCallbackQuery",

            json={

                "callback_query_id": callback["id"]

            },

            timeout=10

        )

        return "OK"

    # =====================

    # NORMAL MESSAGE

    # =====================

    message = data.get("message", {})

    chat = message.get("chat", {})

    user = message.get("from", {})

    chat_id = chat.get("id")

    user_id = user.get("id")

    if not chat_id or not user_id:

        return "OK"

    first_name = user.get("first_name", "Student")

    add_user(user_id, first_name)

    # =====================

    # TEXT

    # =====================

    text = message.get("text", "")

    if text == "/start":

        show_menu(chat_id)

        return "OK"

    # =====================

    # COMMANDS

    # =====================

    if text.startswith("/ppt "):

        topic = text[5:].strip()

        if not topic:

            send_message(chat_id, "Tema ýaz.")

            return "OK"

        if not use_credit(user_id):

            send_message(chat_id, "❌ Balansyň gutardy.")

            return "OK"

        send_message(chat_id, "⏳ Prezentasiýa taýýarlanýar...")

        try:

            file = create_presentation(topic)

            send_document(

                chat_id,

                file,

                "📊 Student Bot tarapyndan döredilen prezentasiýa"

            )

        except Exception as e:

            send_message(chat_id, f"❌ Ýalňyşlyk: {str(e)}")

        return "OK"

    if text.startswith("/referat "):

        topic = text[9:].strip()

        if not use_credit(user_id):

            send_message(chat_id, "❌ Balansyň gutardy.")

            return "OK"

        send_message(chat_id, "⏳ Referat taýýarlanýar...")

        answer = ask_ai(

            f"{topic} barada uniwersitet üçin referat ýaz."

        )

        send_message(chat_id, answer)

        return "OK"

    if text.startswith("/word "):

        topic = text[6:].strip()

        if not use_credit(user_id):

            send_message(chat_id, "❌ Balansyň gutardy.")

            return "OK"

        send_message(chat_id, "⏳ Word faýly taýýarlanýar...")

        try:

            file = create_word(topic)

            send_document(

                chat_id,

                file,

                "📄 Student Bot Word faýly"

            )

        except Exception as e:

            send_message(chat_id, f"❌ Ýalňyşlyk: {str(e)}")

        return "OK"

    if text.startswith("/pdf "):

        topic = text[5:].strip()

        if not use_credit(user_id):

            send_message(chat_id, "❌ Balansyň gutardy.")

            return "OK"

        send_message(chat_id, "⏳ PDF taýýarlanýar...")

        try:

            file = create_pdf(topic)

            send_document(

                chat_id,

                file,

                "📕 Student Bot PDF faýly"

            )

        except Exception as e:

            send_message(chat_id, f"❌ Ýalňyşlyk: {str(e)}")

        return "OK"

    if text.startswith("/translate "):

        content = text[11:].strip()

        if not use_credit(user_id):

            send_message(chat_id, "❌ Balansyň gutardy.")

            return "OK"

        answer = ask_ai(

            f"Şu teksti terjime et:\n\n{content}"

        )

        send_message(chat_id, answer)

        return "OK"

    # =====================

    # PHOTO

    # =====================

    if "photo" in message:

        if not use_credit(user_id):

            send_message(chat_id, "❌ Balansyň gutardy.")

            return "OK"

        send_message(

            chat_id,

            "📷 Suraty okaýaryn we meseläni çözýärin..."

        )

        try:

            photo = message["photo"][-1]

            image_bytes = get_file_bytes(photo["file_id"])

            answer = solve_image(image_bytes)

            send_message(chat_id, answer)

        except Exception as e:

            send_message(

                chat_id,

                f"❌ Surat işlenende ýalňyşlyk: {str(e)}"

            )

        return "OK"

    # =====================

    # NORMAL AI QUESTION

    # =====================

    if text:

        if not use_credit(user_id):

            send_message(

                chat_id,

                "❌ Balansyň gutardy.\n\n"

                "Täze kredit almak üçin administrator bilen habarlaş."

            )

            return "OK"

        try:

            answer = ask_ai(text)

            send_message(chat_id, answer)

        except Exception as e:

            send_message(

                chat_id,

                f"❌ AI ýalňyşlygy: {str(e)}"

            )

    return "OK"

# =========================

# START

# =========================

init_db()

if __name__ == "__main__":

    port = int(os.environ.get("PORT", 10000))

    app.run(

        host="0.0.0.0",

        port=port

    )
